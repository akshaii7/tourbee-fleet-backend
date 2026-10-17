import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta
from decimal import Decimal

from django.http import HttpResponse
from django.db.models import Sum, Count, Q, Case, When
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import permissions, status

from apps.accounts.permissions import IsAdminUserRole
from apps.vehicles.models import Vehicle
from apps.drivers.models import Driver
from apps.submissions.models import EarningsSubmission, ExpenseSubmission, SubmissionStatus
from apps.finance.models import SalaryCalculation, SalaryPayment, RentPayment, BusinessSettings
from apps.maintenance.models import MaintenanceRecord

class DashboardSummaryView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user
        today = timezone.localdate()
        first_of_month = today.replace(day=1)

        if user.is_admin:
            # Fleet Admin Dashboard Metrics - combined counts in single queries
            veh_counts = Vehicle.objects.aggregate(
                total=Count('id'),
                active=Count(Case(When(status=Vehicle.Status.ACTIVE, then=1))),
                maintenance=Count(Case(When(status=Vehicle.Status.MAINTENANCE, then=1))),
            )
            total_vehicles = veh_counts['total'] or 0
            active_vehicles = veh_counts['active'] or 0
            maintenance_vehicles = veh_counts['maintenance'] or 0

            driver_counts = Driver.objects.aggregate(
                total=Count('id'),
                active=Count(Case(When(is_active=True, then=1))),
                profit_sharing=Count(Case(When(is_active=True, driver_type=Driver.DriverType.PROFIT_SHARING, then=1))),
                daily_rent=Count(Case(When(is_active=True, driver_type=Driver.DriverType.DAILY_RENT, then=1))),
            )
            total_drivers = driver_counts['total'] or 0
            active_drivers = driver_counts['active'] or 0
            profit_sharing_drivers = driver_counts['profit_sharing'] or 0
            daily_rent_drivers = driver_counts['daily_rent'] or 0

            pending_earnings = EarningsSubmission.objects.filter(status=SubmissionStatus.PENDING).count()
            pending_expenses = ExpenseSubmission.objects.filter(status=SubmissionStatus.PENDING).count()
            total_pending = pending_earnings + pending_expenses

            # Financial metrics for current month - combined aggregations
            month_earnings_agg = EarningsSubmission.objects.filter(
                status=SubmissionStatus.APPROVED,
                date__gte=first_of_month,
                date__lte=today
            ).aggregate(
                gross=Sum('gross_earnings'),
                comm=Sum('uber_commission'),
                cash=Sum('cash_collected'),
                trips=Sum('trips_count')
            )
            gross_revenue_month = month_earnings_agg['gross'] or Decimal('0.00')
            commission_month = month_earnings_agg['comm'] or Decimal('0.00')
            cash_collected_month = month_earnings_agg['cash'] or Decimal('0.00')
            trips_count_month = month_earnings_agg['trips'] or 0

            operating_expenses_month = ExpenseSubmission.objects.filter(
                status=SubmissionStatus.APPROVED,
                date__gte=first_of_month,
                date__lte=today
            ).aggregate(v=Sum('amount'))['v'] or Decimal('0.00')

            maintenance_cost_month = MaintenanceRecord.objects.filter(
                date__gte=first_of_month,
                date__lte=today
            ).aggregate(v=Sum('cost'))['v'] or Decimal('0.00')

            rent_collected_month = RentPayment.objects.filter(
                payment_date__gte=first_of_month,
                payment_date__lte=today
            ).aggregate(v=Sum('amount_paid'))['v'] or Decimal('0.00')

            # Salary Calculations for the month in single aggregate
            salary_agg = SalaryCalculation.objects.filter(
                period_end__gte=first_of_month
            ).aggregate(
                due=Sum('net_payable_to_driver'),
                share=Sum('company_share_50')
            )
            driver_payouts_due = salary_agg['due'] or Decimal('0.00')
            company_profit_share_month = salary_agg['share'] or Decimal('0.00')

            # Net Company Profit = Company 50% Profit Share + Rent Collected - Maintenance Expenses
            net_company_profit = company_profit_share_month + rent_collected_month - maintenance_cost_month

            # Document Expiry Warnings (Insurance or Fitness <= 15 days or expired)
            expiring_vehicles = []
            for v in Vehicle.objects.only('id', 'registration_number', 'make', 'model', 'insurance_expiry', 'fitness_expiry'):
                ins_status = v.insurance_status
                fit_status = v.fitness_status
                if ins_status in ['EXPIRED', 'EXPIRING_SOON'] or fit_status in ['EXPIRED', 'EXPIRING_SOON']:
                    expiring_vehicles.append({
                        'id': v.id,
                        'registration_number': v.registration_number,
                        'model': f"{v.make} {v.model}",
                        'insurance_expiry': v.insurance_expiry.strftime('%d-%m-%Y') if v.insurance_expiry else None,
                        'insurance_status': ins_status,
                        'fitness_expiry': v.fitness_expiry.strftime('%d-%m-%Y') if v.fitness_expiry else None,
                        'fitness_status': fit_status,
                    })

            # Daily trend for past 14 days - batched in 2 group-by queries instead of 42 queries
            fourteen_days_ago = today - timedelta(days=13)
            earnings_by_day = {
                row['date']: row
                for row in EarningsSubmission.objects.filter(
                    status=SubmissionStatus.APPROVED,
                    date__gte=fourteen_days_ago,
                    date__lte=today
                ).values('date').annotate(
                    gross=Sum('gross_earnings'),
                    comm=Sum('uber_commission')
                )
            }
            expenses_by_day = {
                row['date']: row['total']
                for row in ExpenseSubmission.objects.filter(
                    status=SubmissionStatus.APPROVED,
                    date__gte=fourteen_days_ago,
                    date__lte=today
                ).values('date').annotate(
                    total=Sum('amount')
                )
            }

            daily_stats = []
            for i in range(14):
                day = fourteen_days_ago + timedelta(days=i)
                day_earn = earnings_by_day.get(day, {})
                day_gross = day_earn.get('gross') or Decimal('0.00')
                day_comm = day_earn.get('comm') or Decimal('0.00')
                day_exp = expenses_by_day.get(day) or Decimal('0.00')
                daily_stats.append({
                    'date': day.strftime('%d-%m'),
                    'full_date': day.strftime('%d-%m-%Y'),
                    'gross': float(day_gross),
                    'commission': float(day_comm),
                    'operating_expenses': float(day_exp),
                    'net': float(day_gross - day_comm - day_exp)
                })

            return Response({
                'role': 'ADMIN',
                'overview': {
                    'total_vehicles': total_vehicles,
                    'active_vehicles': active_vehicles,
                    'maintenance_vehicles': maintenance_vehicles,
                    'total_drivers': total_drivers,
                    'active_drivers': active_drivers,
                    'profit_sharing_drivers': profit_sharing_drivers,
                    'daily_rent_drivers': daily_rent_drivers,
                    'pending_approvals': total_pending,
                    'pending_earnings_count': pending_earnings,
                    'pending_expenses_count': pending_expenses,
                },
                'financials_month': {
                    'gross_revenue': float(gross_revenue_month),
                    'uber_commission': float(commission_month),
                    'operating_expenses': float(operating_expenses_month),
                    'maintenance_cost': float(maintenance_cost_month),
                    'rent_collected': float(rent_collected_month),
                    'trips_count': trips_count_month,
                    'company_profit_share': float(company_profit_share_month),
                    'net_company_profit': float(net_company_profit),
                    'cash_collected_by_drivers': float(cash_collected_month),
                },
                'daily_stats': daily_stats,
                'expiring_vehicles': expiring_vehicles,
            })

        else:
            # Driver personal dashboard metrics
            if not hasattr(user, 'driver_profile'):
                return Response({'error': 'Driver profile not linked to this account.'}, status=status.HTTP_404_NOT_FOUND)

            driver = user.driver_profile
            # Submissions stats for driver in single aggregate
            my_earnings_agg = EarningsSubmission.objects.filter(
                driver=driver,
                date__gte=first_of_month,
                date__lte=today,
                status=SubmissionStatus.APPROVED
            ).aggregate(
                gross=Sum('gross_earnings'),
                comm=Sum('uber_commission'),
                cash=Sum('cash_collected'),
                trips=Sum('trips_count')
            )
            my_gross = my_earnings_agg['gross'] or Decimal('0.00')
            my_comm = my_earnings_agg['comm'] or Decimal('0.00')
            my_cash = my_earnings_agg['cash'] or Decimal('0.00')
            my_trips = my_earnings_agg['trips'] or 0

            my_exp_total = ExpenseSubmission.objects.filter(
                driver=driver,
                date__gte=first_of_month,
                date__lte=today,
                status=SubmissionStatus.APPROVED
            ).aggregate(v=Sum('amount'))['v'] or Decimal('0.00')

            pending_count = EarningsSubmission.objects.filter(driver=driver, status=SubmissionStatus.PENDING).count() + \
                            ExpenseSubmission.objects.filter(driver=driver, status=SubmissionStatus.PENDING).count()

            correction_count = EarningsSubmission.objects.filter(driver=driver, status=SubmissionStatus.CORRECTION_REQUESTED).count() + \
                               ExpenseSubmission.objects.filter(driver=driver, status=SubmissionStatus.CORRECTION_REQUESTED).count()

            # Latest salary calculation (if profit sharing)
            latest_calc = None
            if driver.driver_type == Driver.DriverType.PROFIT_SHARING:
                calc_obj = SalaryCalculation.objects.filter(driver=driver).order_by('-period_end').first()
                if calc_obj:
                    latest_calc = {
                        'id': calc_obj.id,
                        'period': f"{calc_obj.period_start.strftime('%d-%m-%Y')} to {calc_obj.period_end.strftime('%d-%m-%Y')}",
                        'distributable_profit': float(calc_obj.distributable_profit),
                        'driver_share_50': float(calc_obj.driver_share_50),
                        'cash_collected_deduction': float(calc_obj.cash_collected_deduction),
                        'reimbursements': float(calc_obj.driver_reimbursements),
                        'net_payable': float(calc_obj.net_payable_to_driver),
                        'status': calc_obj.status,
                    }

            # Rent status (if daily rent)
            rent_info = None
            if driver.driver_type == Driver.DriverType.DAILY_RENT:
                rent_agg = RentPayment.objects.filter(driver=driver).aggregate(
                    due=Sum('amount_due'),
                    paid=Sum('amount_paid')
                )
                due = rent_agg['due'] or Decimal('0.00')
                paid = rent_agg['paid'] or Decimal('0.00')
                rent_info = {
                    'daily_rate': float(driver.daily_rent_amount),
                    'total_due': float(due),
                    'total_paid': float(paid),
                    'outstanding_balance': float(max(Decimal('0.00'), due - paid)),
                }

            # Recent 7 days earnings batched in 1 query instead of 7
            seven_days_ago = today - timedelta(days=6)
            driver_subs_by_day = {
                s.date: s
                for s in EarningsSubmission.objects.filter(
                    driver=driver,
                    date__gte=seven_days_ago,
                    date__lte=today
                )
            }
            recent_daily = []
            for i in range(7):
                d = seven_days_ago + timedelta(days=i)
                sub = driver_subs_by_day.get(d)
                recent_daily.append({
                    'date': d.strftime('%d-%m'),
                    'gross': float(sub.gross_earnings) if sub and sub.status == SubmissionStatus.APPROVED else 0.0,
                    'status': sub.status if sub else 'NO_SUBMISSION'
                })

            return Response({
                'role': 'DRIVER',
                'driver': {
                    'id': driver.id,
                    'name': driver.full_name,
                    'driver_type': driver.driver_type,
                    'assigned_vehicle': driver.assigned_vehicle.registration_number if driver.assigned_vehicle else 'None',
                    'assigned_vehicle_name': f"{driver.assigned_vehicle.make} {driver.assigned_vehicle.model}" if driver.assigned_vehicle else None,
                },
                'month_summary': {
                    'gross_earnings': float(my_gross),
                    'uber_commission': float(my_comm),
                    'cash_collected': float(my_cash),
                    'trips_count': my_trips,
                    'operating_expenses': float(my_exp_total),
                },
                'pending_submissions_count': pending_count,
                'correction_requested_count': correction_count,
                'latest_salary_calculation': latest_calc,
                'rent_info': rent_info,
                'recent_daily': recent_daily
            })

class FinancialReportsView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if not (getattr(request.user, 'is_admin', False) or getattr(request.user, 'role', '') == 'ADMIN' or request.user.is_staff):
            return Response({'error': 'Admin privileges required'}, status=status.HTTP_403_FORBIDDEN)

        today = timezone.localdate()
        raw_start = request.query_params.get('start_date')
        raw_end = request.query_params.get('end_date')

        # Allow empty or 'all' to show all-time data
        start_date_str = raw_start.strip() if raw_start else ''
        end_date_str = raw_end.strip() if raw_end else ''

        # If completely omitted from query params, default to current month
        if raw_start is None and raw_end is None:
            start_date_str = today.replace(day=1).strftime('%Y-%m-%d')
            end_date_str = today.strftime('%Y-%m-%d')

        from fleet.models import (
            DailyCalculation, DriverProfile, Vehicle as FleetVehicle,
            SalaryPayment as FleetSalaryPayment, PaymentModeChoices
        )

        # 1. Fleet Daily Calculations - filtered safely
        calcs_qs = DailyCalculation.objects.all()
        if start_date_str and start_date_str.lower() != 'all':
            try:
                calcs_qs = calcs_qs.filter(date__gte=start_date_str)
            except Exception:
                pass
        if end_date_str and end_date_str.lower() != 'all':
            try:
                calcs_qs = calcs_qs.filter(date__lte=end_date_str)
            except Exception:
                pass

        calcs_agg = calcs_qs.aggregate(
            fleet_gross_uber=Sum(
                Case(When(payment_mode=PaymentModeChoices.REVENUE_SHARING, then='A'), default=Decimal('0.00'))
            ),
            fleet_cash=Sum('B'),
            fleet_expenses=Sum('D'),
            fleet_cng=Sum('E'),
            fleet_toll=Sum('F'),
            fleet_driver_share=Sum(
                Case(When(payment_mode=PaymentModeChoices.REVENUE_SHARING, then='driver_share'), default=Decimal('0.00'))
            ),
            fleet_company_share=Sum(
                Case(When(payment_mode=PaymentModeChoices.REVENUE_SHARING, then='company_share'), default=Decimal('0.00'))
            ),
            fleet_rent=Sum(
                Case(When(payment_mode=PaymentModeChoices.FIXED_RENT, then='J'), default=Decimal('0.00'))
            ),
            fleet_trips=Count('id'),
        )

        fleet_gross_uber = calcs_agg['fleet_gross_uber'] or Decimal('0.00')
        fleet_cash = calcs_agg['fleet_cash'] or Decimal('0.00')
        fleet_expenses = calcs_agg['fleet_expenses'] or Decimal('0.00')
        fleet_cng = calcs_agg['fleet_cng'] or Decimal('0.00')
        fleet_toll = calcs_agg['fleet_toll'] or Decimal('0.00')
        fleet_driver_share = calcs_agg['fleet_driver_share'] or Decimal('0.00')
        fleet_company_share = calcs_agg['fleet_company_share'] or Decimal('0.00')
        fleet_rent = calcs_agg['fleet_rent'] or Decimal('0.00')
        fleet_trips = calcs_agg['fleet_trips'] or 0

        # 2. Legacy Earnings / Expenses / Rent in single aggregated queries
        legacy_earnings = EarningsSubmission.objects.filter(status=SubmissionStatus.APPROVED)
        if start_date_str and start_date_str.lower() != 'all':
            try:
                legacy_earnings = legacy_earnings.filter(date__gte=start_date_str)
            except Exception:
                pass
        if end_date_str and end_date_str.lower() != 'all':
            try:
                legacy_earnings = legacy_earnings.filter(date__lte=end_date_str)
            except Exception:
                pass

        l_earn_agg = legacy_earnings.aggregate(
            gross=Sum('gross_earnings'),
            comm=Sum('uber_commission'),
            cash=Sum('cash_collected'),
            trips=Sum('trips_count')
        )
        legacy_gross = l_earn_agg['gross'] or Decimal('0.00')
        legacy_comm = l_earn_agg['comm'] or Decimal('0.00')
        legacy_cash = l_earn_agg['cash'] or Decimal('0.00')
        legacy_trips = l_earn_agg['trips'] or 0

        legacy_expenses = ExpenseSubmission.objects.filter(status=SubmissionStatus.APPROVED)
        if start_date_str and start_date_str.lower() != 'all':
            try:
                legacy_expenses = legacy_expenses.filter(date__gte=start_date_str)
            except Exception:
                pass
        if end_date_str and end_date_str.lower() != 'all':
            try:
                legacy_expenses = legacy_expenses.filter(date__lte=end_date_str)
            except Exception:
                pass

        l_exp_agg = legacy_expenses.aggregate(
            op_exp=Sum('amount'),
            cng=Sum(Case(When(category__code='CNG', then='amount'), default=Decimal('0.00'))),
            toll=Sum(Case(When(category__code='TOLL', then='amount'), default=Decimal('0.00'))),
        )
        legacy_op_exp = l_exp_agg['op_exp'] or Decimal('0.00')
        legacy_cng = l_exp_agg['cng'] or Decimal('0.00')
        legacy_toll = l_exp_agg['toll'] or Decimal('0.00')

        legacy_rent_qs = RentPayment.objects.all()
        if start_date_str and start_date_str.lower() != 'all':
            try:
                legacy_rent_qs = legacy_rent_qs.filter(payment_date__gte=start_date_str)
            except Exception:
                pass
        if end_date_str and end_date_str.lower() != 'all':
            try:
                legacy_rent_qs = legacy_rent_qs.filter(payment_date__lte=end_date_str)
            except Exception:
                pass
        legacy_rent = legacy_rent_qs.aggregate(v=Sum('amount_paid'))['v'] or Decimal('0.00')

        # 3. Maintenance Records
        maint_qs = MaintenanceRecord.objects.all()
        if start_date_str and start_date_str.lower() != 'all':
            try:
                maint_qs = maint_qs.filter(date__gte=start_date_str)
            except Exception:
                pass
        if end_date_str and end_date_str.lower() != 'all':
            try:
                maint_qs = maint_qs.filter(date__lte=end_date_str)
            except Exception:
                pass
        maintenance_costs = maint_qs.aggregate(v=Sum('cost'))['v'] or Decimal('0.00')

        # Combined totals
        gross_uber = fleet_gross_uber + legacy_gross
        uber_comm = legacy_comm
        cash_collected = fleet_cash + legacy_cash
        total_trips = fleet_trips + legacy_trips
        cng_expenses = fleet_cng + legacy_cng
        toll_expenses = fleet_toll + legacy_toll
        operating_expenses = fleet_expenses + legacy_op_exp
        other_expenses = max(Decimal('0.00'), operating_expenses - cng_expenses - toll_expenses)
        rent_revenue = fleet_rent + legacy_rent
        driver_profit_share = fleet_driver_share
        company_profit_share = fleet_company_share
        company_net_profit = company_profit_share + rent_revenue - maintenance_costs

        # Vehicle Breakdown - batched pre-aggregation to eliminate N+1 queries
        vehicle_breakdown = []
        seen_regs = set()
        all_vehicles = []

        fleet_vehicles_by_reg = {
            fv.registration_number: fv.id
            for fv in FleetVehicle.objects.only('id', 'registration_number')
        }

        for v in Vehicle.objects.only('id', 'registration_number', 'make', 'model'):
            if v.registration_number not in seen_regs:
                seen_regs.add(v.registration_number)
                all_vehicles.append({
                    'id': v.id,
                    'registration': v.registration_number,
                    'model': f"{v.make} {v.model}",
                    'apps_vehicle_id': v.id,
                    'fleet_vehicle_id': fleet_vehicles_by_reg.get(v.registration_number)
                })
        for fv in FleetVehicle.objects.only('id', 'registration_number', 'make', 'model'):
            if fv.registration_number not in seen_regs:
                seen_regs.add(fv.registration_number)
                all_vehicles.append({
                    'id': fv.id,
                    'registration': fv.registration_number,
                    'model': f"{fv.make} {fv.model}",
                    'apps_vehicle_id': None,
                    'fleet_vehicle_id': fv.id
                })

        calcs_by_veh_reg = {}
        calcs_by_veh_id = {}
        for row in calcs_qs.filter(vehicle__isnull=False).values('vehicle__registration_number', 'vehicle_id').annotate(
            f_gross=Sum('A'),
            f_exp=Sum('D'),
            f_cng=Sum('E'),
            f_trips=Count('id')
        ):
            if row['vehicle__registration_number']:
                calcs_by_veh_reg[row['vehicle__registration_number']] = row
            if row['vehicle_id']:
                calcs_by_veh_id[row['vehicle_id']] = row

        legacy_earn_by_veh = {
            row['vehicle_id']: row
            for row in legacy_earnings.filter(vehicle_id__isnull=False).values('vehicle_id').annotate(
                gross=Sum('gross_earnings'),
                trips=Sum('trips_count')
            )
        }
        legacy_exp_by_veh = {
            row['vehicle_id']: row['tot']
            for row in legacy_expenses.filter(vehicle_id__isnull=False).values('vehicle_id').annotate(
                tot=Sum('amount')
            )
        }
        legacy_maint_by_veh = {
            row['vehicle_id']: row['tot']
            for row in maint_qs.filter(vehicle_id__isnull=False).values('vehicle_id').annotate(
                tot=Sum('cost')
            )
        }

        for vinfo in all_vehicles:
            reg = vinfo['registration']
            fv_id = vinfo['fleet_vehicle_id']
            row_c = calcs_by_veh_reg.get(reg) or (calcs_by_veh_id.get(fv_id) if fv_id else None) or {}
            f_gross = row_c.get('f_gross') or Decimal('0.00')
            f_exp = row_c.get('f_exp') or Decimal('0.00')
            f_cng = row_c.get('f_cng') or Decimal('0.00')
            f_trips = row_c.get('f_trips') or 0

            l_gross = Decimal('0.00')
            l_exp = Decimal('0.00')
            l_maint = Decimal('0.00')
            l_trips = 0
            if vinfo['apps_vehicle_id']:
                app_v_id = vinfo['apps_vehicle_id']
                l_earn = legacy_earn_by_veh.get(app_v_id, {})
                l_gross = l_earn.get('gross') or Decimal('0.00')
                l_trips = l_earn.get('trips') or 0
                l_exp = legacy_exp_by_veh.get(app_v_id) or Decimal('0.00')
                l_maint = legacy_maint_by_veh.get(app_v_id) or Decimal('0.00')

            tot_gross = f_gross + l_gross
            tot_exp = f_exp + f_cng + l_exp
            tot_trips = f_trips + l_trips

            vehicle_breakdown.append({
                'id': vinfo['id'],
                'registration': reg,
                'model': vinfo['model'],
                'gross_earnings': float(tot_gross),
                'operating_expenses': float(tot_exp),
                'maintenance_cost': float(l_maint),
                'trips': tot_trips,
                'net_fleet_contribution': float(tot_gross - tot_exp - l_maint)
            })

        # Driver Breakdown - batched pre-aggregation to eliminate N+1 queries
        driver_breakdown = []
        calcs_by_user_id = {}
        calcs_by_driver_id = {}
        for row in calcs_qs.values('driver__user_id', 'driver_id').annotate(
            d_gross=Sum('A'),
            d_exp=Sum('D'),
            d_cng=Sum('E'),
            d_share=Sum('driver_share'),
            d_rent=Sum('J'),
            d_trips=Count('id')
        ):
            if row['driver__user_id']:
                calcs_by_user_id[row['driver__user_id']] = row
            if row['driver_id']:
                calcs_by_driver_id[row['driver_id']] = row

        legacy_earn_by_driver = {
            row['driver_id']: row
            for row in legacy_earnings.filter(driver_id__isnull=False).values('driver_id').annotate(
                gross=Sum('gross_earnings'),
                comm=Sum('uber_commission'),
                trips=Sum('trips_count')
            )
        }
        legacy_exp_by_driver = {
            row['driver_id']: row['tot']
            for row in legacy_expenses.filter(driver_id__isnull=False).values('driver_id').annotate(
                tot=Sum('amount')
            )
        }

        seen_user_ids = set()
        for d in Driver.objects.select_related('user', 'assigned_vehicle'):
            d_user_id = d.user_id
            if d_user_id:
                seen_user_ids.add(d_user_id)
            row_c = calcs_by_user_id.get(d_user_id) or calcs_by_driver_id.get(d.id) or {}
            d_gross = row_c.get('d_gross') or Decimal('0.00')
            d_exp = row_c.get('d_exp') or Decimal('0.00')
            d_cng = row_c.get('d_cng') or Decimal('0.00')
            d_share = row_c.get('d_share') or Decimal('0.00')
            d_rent = row_c.get('d_rent') or Decimal('0.00')
            d_trips = row_c.get('d_trips') or 0

            l_earn = legacy_earn_by_driver.get(d.id, {})
            l_gross = l_earn.get('gross') or Decimal('0.00')
            l_comm = l_earn.get('comm') or Decimal('0.00')
            l_trips = l_earn.get('trips') or 0
            l_exp = legacy_exp_by_driver.get(d.id) or Decimal('0.00')

            tot_gross = d_gross + l_gross
            tot_exp = d_exp + d_cng + l_exp
            tot_trips = d_trips + l_trips

            driver_breakdown.append({
                'id': d.id,
                'name': d.full_name,
                'driver_type': d.driver_type,
                'vehicle': d.assigned_vehicle.registration_number if d.assigned_vehicle else 'None',
                'gross_earnings': float(tot_gross),
                'commission': float(l_comm),
                'expenses': float(tot_exp),
                'driver_salary': float(d_share),
                'rent_due': float(d_rent),
                'trips': tot_trips,
            })

        for dp in DriverProfile.objects.select_related('user', 'assigned_vehicle'):
            if dp.user_id not in seen_user_ids:
                seen_user_ids.add(dp.user_id)
                row_c = calcs_by_user_id.get(dp.user_id) or calcs_by_driver_id.get(dp.id) or {}
                d_gross = row_c.get('d_gross') or Decimal('0.00')
                d_exp = row_c.get('d_exp') or Decimal('0.00')
                d_cng = row_c.get('d_cng') or Decimal('0.00')
                d_share = row_c.get('d_share') or Decimal('0.00')
                d_rent = row_c.get('d_rent') or Decimal('0.00')
                d_trips = row_c.get('d_trips') or 0
                tot_gross = d_gross
                tot_exp = d_exp + d_cng
                driver_type_name = 'DAILY_RENT' if dp.payment_mode == PaymentModeChoices.FIXED_RENT else 'PROFIT_SHARING'
                driver_breakdown.append({
                    'id': dp.id,
                    'name': dp.user.get_full_name() or dp.user.username,
                    'driver_type': driver_type_name,
                    'vehicle': dp.assigned_vehicle.registration_number if dp.assigned_vehicle else 'None',
                    'gross_earnings': float(tot_gross),
                    'commission': 0.0,
                    'expenses': float(tot_exp),
                    'driver_salary': float(d_share),
                    'rent_due': float(d_rent),
                    'trips': d_trips,
                })

        return Response({
            'period': {'start_date': start_date_str, 'end_date': end_date_str},
            'summary': {
                'gross_uber_earnings': float(gross_uber),
                'uber_commission': float(uber_comm),
                'cash_collected': float(cash_collected),
                'total_trips': total_trips,
                'operating_expenses': float(operating_expenses),
                'cng_expenses': float(cng_expenses),
                'toll_expenses': float(toll_expenses),
                'other_expenses': float(other_expenses),
                'rent_revenue': float(rent_revenue),
                'maintenance_costs': float(maintenance_costs),
                'driver_profit_share': float(driver_profit_share),
                'company_profit_share': float(company_profit_share),
                'company_net_profit': float(company_net_profit),
            },
            'vehicle_breakdown': vehicle_breakdown,
            'driver_breakdown': driver_breakdown,
        })


class ExportDataView(APIView):
    permission_classes = [permissions.AllowAny]

    def perform_content_negotiation(self, request, force=False):
        from rest_framework.renderers import JSONRenderer
        return (JSONRenderer(), 'application/json')

    def get(self, request):
        user = request.user
        if not (user and user.is_authenticated):
            token_str = request.query_params.get('token')
            if token_str:
                from rest_framework_simplejwt.tokens import AccessToken
                from django.contrib.auth import get_user_model
                try:
                    access_token = AccessToken(token_str)
                    user_id = access_token['user_id']
                    User = get_user_model()
                    user = User.objects.get(id=user_id)
                except Exception:
                    return Response({'error': 'Invalid or expired token.'}, status=status.HTTP_401_UNAUTHORIZED)
            else:
                return Response({'error': 'Authentication credentials were not provided.'}, status=status.HTTP_401_UNAUTHORIZED)

        if not (getattr(user, 'is_admin', False) or getattr(user, 'role', '') == 'ADMIN' or user.is_staff):
            return Response({'error': 'Admin permissions required.'}, status=status.HTTP_403_FORBIDDEN)

        report_type = request.query_params.get('type', 'earnings')
        export_format = request.query_params.get('format', 'csv').lower()
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')

        from fleet.models import DailyCalculation

        if report_type == 'earnings':
            qs = DailyCalculation.objects.all().select_related('driver__user', 'vehicle').order_by('-date', '-created_at')
            if start_date:
                qs = qs.filter(date__gte=start_date)
            if end_date:
                qs = qs.filter(date__lte=end_date)

            headers = [
                'ID', 'Date', 'Driver Name', 'Phone', 'Vehicle Reg', 'Payment Mode',
                'Working Shift', 'Gross Uber Earnings (INR)', 'Passenger Cash (INR)',
                'Operating Expenses (INR)', 'CNG Fuel Bill (INR)', 'Toll / Refund (INR)',
                'Profit Amount (INR)', 'Driver Salary 50% (INR)', 'TourBee Share 50% (INR)',
                'Fixed Rent (INR)', 'Payment Status'
            ]
            rows = []
            for r in qs:
                d_name = r.driver.user.get_full_name() or r.driver.user.username if r.driver else 'N/A'
                d_phone = r.driver.phone if r.driver else ''
                v_reg = r.vehicle.registration_number if r.vehicle else 'N/A'
                rows.append([
                    r.id,
                    r.date.strftime('%d-%m-%Y'),
                    d_name,
                    d_phone,
                    v_reg,
                    r.get_payment_mode_display(),
                    r.get_working_day_type_display(),
                    float(r.A),
                    float(r.B),
                    float(r.D),
                    float(r.E),
                    float(r.F),
                    float(r.G),
                    float(r.driver_share),
                    float(r.company_share),
                    float(r.J),
                    r.status
                ])

            # Also append legacy EarningsSubmission if any
            legacy_qs = EarningsSubmission.objects.all().select_related('driver__user', 'vehicle').order_by('-date')
            if start_date:
                legacy_qs = legacy_qs.filter(date__gte=start_date)
            if end_date:
                legacy_qs = legacy_qs.filter(date__lte=end_date)
            for r in legacy_qs:
                rows.append([
                    f"LEGACY-{r.id}",
                    r.date.strftime('%d-%m-%Y'),
                    r.driver.full_name,
                    getattr(r.driver.user, 'phone_number', ''),
                    r.vehicle.registration_number,
                    r.driver.driver_type,
                    'FULL_DAY',
                    float(r.gross_earnings),
                    float(r.cash_collected),
                    float(r.uber_commission),
                    0.0,
                    0.0,
                    float(r.net_earnings),
                    float(r.net_earnings) * 0.5,
                    float(r.net_earnings) * 0.5,
                    0.0,
                    r.status
                ])

            filename = f"tourbee_earnings_{timezone.localdate().strftime('%d_%m_%Y')}"

        elif report_type == 'expenses':
            qs = ExpenseSubmission.objects.all().select_related('driver__user', 'vehicle', 'category').order_by('-date')
            if start_date:
                qs = qs.filter(date__gte=start_date)
            if end_date:
                qs = qs.filter(date__lte=end_date)

            headers = [
                'ID', 'Date', 'Driver Name', 'Vehicle Reg', 'Category',
                'Amount (INR)', 'Vendor', 'Odometer', 'Fuel Qty', 'Paid By Driver', 'Status'
            ]
            rows = []
            for r in qs:
                rows.append([
                    r.id,
                    r.date.strftime('%d-%m-%Y'),
                    r.driver.full_name,
                    r.vehicle.registration_number,
                    r.category.name,
                    float(r.amount),
                    r.vendor_name or '',
                    r.odometer_reading or '',
                    float(r.fuel_quantity_kg_or_l or 0),
                    'Yes' if r.paid_by_driver else 'No',
                    r.status
                ])

            # Also add CNG and Operating Expenses from DailyCalculation
            calcs_qs = DailyCalculation.objects.filter(Q(D__gt=0) | Q(E__gt=0)).select_related('driver__user', 'vehicle').order_by('-date')
            if start_date:
                calcs_qs = calcs_qs.filter(date__gte=start_date)
            if end_date:
                calcs_qs = calcs_qs.filter(date__lte=end_date)
            for r in calcs_qs:
                d_name = r.driver.user.get_full_name() or r.driver.user.username if r.driver else 'N/A'
                v_reg = r.vehicle.registration_number if r.vehicle else 'N/A'
                if r.E > 0:
                    rows.append([
                        f"CALC-CNG-{r.id}",
                        r.date.strftime('%d-%m-%Y'),
                        d_name,
                        v_reg,
                        'CNG Fuel',
                        float(r.E),
                        'CNG Station',
                        '',
                        0,
                        'Yes',
                        r.status
                    ])
                if r.D > 0:
                    rows.append([
                        f"CALC-EXP-{r.id}",
                        r.date.strftime('%d-%m-%Y'),
                        d_name,
                        v_reg,
                        'Daily Expense',
                        float(r.D),
                        'General',
                        '',
                        0,
                        'Yes',
                        r.status
                    ])

            filename = f"tourbee_expenses_{timezone.localdate().strftime('%d_%m_%Y')}"

        elif report_type == 'maintenance':
            qs = MaintenanceRecord.objects.all().select_related('vehicle').order_by('-date')
            if start_date:
                qs = qs.filter(date__gte=start_date)
            if end_date:
                qs = qs.filter(date__lte=end_date)

            headers = [
                'ID', 'Date', 'Vehicle Reg', 'Service Type', 'Cost (INR)',
                'Garage / Workshop', 'Invoice No', 'Odometer', 'Notes'
            ]
            rows = []
            for r in qs:
                rows.append([
                    r.id,
                    r.date.strftime('%d-%m-%Y'),
                    r.vehicle.registration_number,
                    r.get_maintenance_type_display(),
                    float(r.cost),
                    r.vendor_garage,
                    r.invoice_number or '',
                    r.odometer_reading or '',
                    r.notes or ''
                ])
            filename = f"tourbee_maintenance_{timezone.localdate().strftime('%d_%m_%Y')}"

        else:
            return Response({'error': f'Unsupported report type: {report_type}'}, status=status.HTTP_400_BAD_REQUEST)

        # Output format
        if export_format == 'excel':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = report_type.capitalize()

            header_fill = PatternFill(start_color='064E3B', end_color='064E3B', fill_type='solid')
            header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')

            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            for row in rows:
                ws.append(row)

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
            return response

        else: # CSV
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
            writer = csv.writer(response)
            writer.writerow(headers)
            for row in rows:
                writer.writerow(row)
            return response

